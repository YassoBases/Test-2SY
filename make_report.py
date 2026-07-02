from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()

# Page margins
for section in doc.sections:
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

def set_rtl(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    bidi = OxmlElement('w:bidi')
    pPr.insert(0, bidi)

def heading(doc, text, level=1, color=(30, 100, 200)):
    p = doc.add_paragraph()
    set_rtl(p)
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(16 if level == 1 else 13)
    run.font.color.rgb = RGBColor(*color)
    run.font.name = 'Arial'
    return p

def body(doc, text, bold=False):
    p = doc.add_paragraph()
    set_rtl(p)
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(11)
    run.font.name = 'Arial'
    return p

def bullet(doc, text):
    p = doc.add_paragraph(style='List Bullet')
    set_rtl(p)
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run(text)
    run.font.size = Pt(11)
    run.font.name = 'Arial'
    return p

# Title
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_rtl(p)
run = p.add_run('ملخص الإضافات على مشروع EduSpark')
run.bold = True
run.font.size = Pt(20)
run.font.color.rgb = RGBColor(30, 100, 200)
run.font.name = 'Arial'

body(doc, 'تم تطوير ثلاث أفكار رئيسية على نظام التقويم الذكي في المشروع.')
doc.add_paragraph()

# ═══ Feature 1 ═══
heading(doc, 'الفكرة الأولى: ربط نتائج الاختبارات بالتقويم تلقائياً', 1, (220, 50, 50))
body(doc, 'المشكلة:', bold=True)
body(doc, 'الطالب يحل اختباراً وتُحفظ نتيجته، لكن التقويم الذكي ما يعرف عنها شيئاً. المواد الضعيفة كانت يدوية.')
body(doc, 'شو عملنا:', bold=True)
bullet(doc, 'نتيجة أقل من 60% ← المادة تُضاف تلقائياً للمواد الضعيفة في التقويم + تنبيه لولي الأمر')
bullet(doc, 'نتيجة 80% وأكثر ← المادة تُشال من الضعيفة تلقائياً + إشعار تحسّن')
bullet(doc, 'بطاقة المواد الضعيفة صارت قابلة للضغط وتعرض أسماء المواد')
doc.add_paragraph()

# ═══ Feature 2 ═══
heading(doc, 'الفكرة الثانية: برنامج دراسي شخصي ذكي', 1, (30, 150, 80))
body(doc, 'المشكلة:', bold=True)
body(doc, 'الجدول كان ثابتاً — نفس الجلسة كل يوم، لا يعرف الفرق بين يوم دوام وعطلة.')
body(doc, 'شو عملنا:', bold=True)
body(doc, '1. جدول يختلف حسب اليوم:', bold=True)
bullet(doc, 'أيام الدراسة ← جلسات 40 دقيقة')
bullet(doc, 'أيام العطلة (جمعة/سبت) ← جلسات 75 دقيقة مع شارة مميزة')
body(doc, '2. تأكيد الإنجاز اليومي:', bold=True)
bullet(doc, '✅ أنجزت ← تُسجَّل مكتملة + رسالة تأكيد')
bullet(doc, '❌ لم أنجز ← تُعاد جدولتها لبكرا تلقائياً')
bullet(doc, '↩️ تراجع ← تعود الجلسة لـ "مخططة"')
body(doc, '3. تقرير أسبوعي ذكي:', bold=True)
bullet(doc, 'نسبة الإنجاز هذا الأسبوع مع دائرة مئوية')
bullet(doc, 'عدد الجلسات المكتملة / الفائتة / المتبقية')
bullet(doc, 'ملاحظات ذكية بناءً على الأداء الفعلي')
bullet(doc, 'زر "أعد توليد الجدول" مرتبط فعلياً بالـ API')
doc.add_paragraph()

# ═══ Feature 3 ═══
heading(doc, 'الفكرة الثالثة: نظام التحفيز والإشعارات الذكية', 1, (150, 50, 200))
body(doc, 'المشكلة:', bold=True)
body(doc, 'ما في شيء يحفّز الطالب يفتح التطبيق كل يوم أو يُعلمه بأهمية الوضع.')
body(doc, 'شو عملنا:', bold=True)
body(doc, '1. عداد الأيام المتتالية (Streak):', bold=True)
bullet(doc, 'كل يوم ينجز فيه جلسة يزيد العداد 🔥')
bullet(doc, 'لو ما درس يوم يرجع للصفر')
bullet(doc, 'رسائل تحفيزية: "3 أيام — استمر!" / "أسبوع كامل — أسطورة! 🏆"')
body(doc, '2. إشعارات ذكية:', bold=True)
bullet(doc, '"عندك 2 جلسة اليوم لم تُنجز — كيمياء، رياضيات"')
bullet(doc, '"امتحان كيمياء بعد يومين — راجع الآن!"')
bullet(doc, '"فوّتت 3 جلسات هذا الأسبوع — عدّل وقت جلساتك"')
body(doc, '3. وضع ما قبل الامتحان:', bold=True)
bullet(doc, 'لما يضيف الطالب امتحاناً خلال 3 أيام، الجدول يتكثّف تلقائياً')
bullet(doc, 'جلسات أطول وأكثر تركيزاً على مادة الامتحان')
doc.add_paragraph()

# ═══ UI ═══
heading(doc, 'تحسينات الواجهة', 2, (100, 100, 100))
bullet(doc, 'تبويبات في الشريط الجانبي: تحليل / سلسلة / إشعارات / تقرير')
bullet(doc, 'تبويبات للجدول: التقويم الأسبوعي / قائمة الجلسات')
bullet(doc, 'أنيميشن Count-Up للأرقام في البطاقات العليا')
bullet(doc, 'Stagger animation للجلسات تظهر بشكل متسلسل')
bullet(doc, 'رسائل تأكيد فورية عند كل فعل (إنجاز / غياب / تراجع)')
doc.add_paragraph()

# ═══ Parent ═══
heading(doc, 'ربط ولي الأمر', 2, (100, 100, 100))
bullet(doc, 'قائمة المواد الضعيفة من نتائج الكويز تظهر لولي الأمر')
bullet(doc, 'تنبيهات الامتحانات القادمة مع تحديد إذا المادة ضعيفة')
bullet(doc, 'ملخص ذكي عن وضع الطالب هذا الأسبوع')

doc.save(r'C:\Users\LEGION\Desktop\EduSpark_Features.docx')
print('Done! File saved to Desktop.')
